"""Role-by-role tests for the audit planning API.

Covers the universe register, the re-audit due list, the plan submit/approve
workflow (including the ownership gate added so an auditor cannot submit a
colleague's plan under their own name), engagement creation and status
transitions (including the guard that refuses completion while findings are
unresolved), and the auditee read scoping on the engagement calendar.
"""
import datetime

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone

from apps.accounts.models import AuditTrail, Department, Role
from apps.audit_planning.models import (
    AuditEngagement, AuditPlan, AuditTeamMember, AuditUniverse, Project,
)
from apps.audit_planning.views import BLOCKS_COMPLETION
from apps.common.role_fixtures import (
    RoleFixtureMixin, make_engagement, make_finding, make_plan, make_universe,
    notification_titles, next_seq,
)
from apps.notifications.models import Notification

UNIVERSE_URL = '/api/planning/universe/'
PROJECTS_URL = '/api/planning/projects/'
PLANS_URL = '/api/planning/plans/'
ENGAGEMENTS_URL = '/api/planning/engagements/'

# An engagement is where the audit gets its bounds, so the API refuses one
# without both an objective and a scope. Every test that posts an engagement
# needs the pair; keeping it here means they say so once rather than each
# restating the same two sentences. Tests that *don't* want them (the
# requiredness test at the bottom of AuditEngagementTest) omit them explicitly.
ENGAGEMENT_SCOPE = {
    'objectives': 'Assess the adequacy and effectiveness of the controls under review.',
    'scope': 'The processes, systems and records named in the annual audit plan.',
}


class AuditUniverseRoleAccessTest(RoleFixtureMixin, TestCase):
    """WRITE_AUDIT gates the register; reads stay open to every role."""

    def test_every_role_can_list(self):
        make_universe(department=self.department)
        self.assert_status_by_role(
            {role: 200 for role in self.users},
            lambda client, role: client.get(UNIVERSE_URL),
        )

    def test_only_write_audit_roles_can_create(self):
        def create(client, role):
            return client.post(UNIVERSE_URL, {
                'name': f'Process for {role}',
                # Unique per role: five roles POSTing the same code would fail
                # on the unique constraint rather than on the permission.
                'code': f'UNV-{role[:3].upper()}-{next_seq()}',
                'category': 'process',
                'risk_score': 3.5,
                'audit_frequency': 'Annually',
                'owner': 'Process Owner',
                'department': self.department.id,
                'status': 'active',
            })

        self.assert_status_by_role({
            Role.ADMIN: 201,
            Role.AUDIT_MANAGER: 201,
            Role.SUPERVISOR: 201,
            Role.AUDITOR: 201,
            Role.AUDITEE: 403,
        }, create)

    def test_only_write_audit_roles_can_update(self):
        entry = make_universe(department=self.department)
        self.assert_status_by_role({
            Role.ADMIN: 200,
            Role.AUDIT_MANAGER: 200,
            Role.SUPERVISOR: 200,
            Role.AUDITOR: 200,
            Role.AUDITEE: 403,
        }, lambda client, role: client.patch(
            f'{UNIVERSE_URL}{entry.id}/', {'owner': role}, format='json',
        ))

    def test_auditee_cannot_delete(self):
        entry = make_universe(department=self.department)
        response = self.as_user(self.auditee).delete(f'{UNIVERSE_URL}{entry.id}/')
        self.assertEqual(response.status_code, 403)
        self.assertTrue(AuditUniverse.objects.filter(pk=entry.pk).exists())

    def test_crud_is_audit_logged(self):
        client = self.as_user(self.auditor)
        created = client.post(UNIVERSE_URL, {
            'name': 'Logged Process', 'code': f'UNV-LOG-{next_seq()}',
            'category': 'process', 'risk_score': 2, 'audit_frequency': 'Annually',
            'status': 'active',
        })
        entry_id = created.data['id']
        client.patch(f'{UNIVERSE_URL}{entry_id}/', {'owner': 'New Owner'}, format='json')
        client.delete(f'{UNIVERSE_URL}{entry_id}/')

        actions = list(
            AuditTrail.objects
            .filter(model_name='AuditUniverse', object_id=str(entry_id))
            .values_list('action', flat=True)
        )
        self.assertCountEqual(actions, ['CREATE', 'UPDATE', 'DELETE'])


class AuditUniverseTransferTest(RoleFixtureMixin, TestCase):
    """Bulk import/export of the universe register — Excel (.xlsx) and CSV.

    Export streams the whole register in the import's column layout; import
    upserts by the unique ``code`` and commits valid rows while reporting the
    invalid ones with their spreadsheet row numbers.
    """

    EXPORT_URL = UNIVERSE_URL + 'export/'
    IMPORT_URL = UNIVERSE_URL + 'import/'
    # Must match the canonical header order the exporter writes.
    CSV_HEADER = ('code,name,category,department_code,directorate_code,description,'
                  'owner,risk_score,audit_frequency,last_audited,status')

    def make_csv(self, rows, filename='universe.csv'):
        """Build an uploaded CSV from a list of row lists (each cell stringified)."""
        text = '\n'.join(
            ','.join('' if cell is None else str(cell) for cell in row)
            for row in rows
        )
        return SimpleUploadedFile(
            filename, text.encode('utf-8'), content_type='text/csv',
        )

    @staticmethod
    def make_xlsx(rows, filename='universe.xlsx'):
        from io import BytesIO

        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return SimpleUploadedFile(
            filename, buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def setUp(self):
        super().setUp()
        self.entry = make_universe(
            department=self.department, name='Existing Process', code=f'UNV-X-{next_seq()}',
        )

    def test_export_csv_round_trips_the_registry(self):
        response = self.as_user(self.manager).get(self.EXPORT_URL, {'filetype': 'csv'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/csv')
        self.assertIn('attachment', response['Content-Disposition'])
        # The exporter writes a UTF-8 BOM so Excel opens Amharic text correctly.
        text = response.content.decode('utf-8-sig')
        self.assertTrue(text.startswith(self.CSV_HEADER))
        self.assertIn(self.entry.code, text)
        self.assertIn(self.entry.name, text)
        self.assertIn(self.department.code, text)

    def test_export_xlsx_contains_the_rows(self):
        from io import BytesIO

        import openpyxl
        response = self.as_user(self.manager).get(self.EXPORT_URL, {'filetype': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        workbook = openpyxl.load_workbook(BytesIO(response.content), read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        codes = [str(r[0]) for r in rows if r[0] is not None]
        self.assertIn(self.entry.code, codes)

    def test_export_rejects_unknown_format(self):
        response = self.as_user(self.manager).get(self.EXPORT_URL, {'filetype': 'pdf'})
        self.assertEqual(response.status_code, 400)

    def test_import_creates_new_rows_from_csv(self):
        code = f'UNV-IMP-{next_seq()}'
        response = self.as_user(self.auditor).post(self.IMPORT_URL, {
            'file': self.make_csv([
                self.CSV_HEADER.split(','),
                [code, 'Imported Process', 'process', self.department.code, '', '',
                 '', '3.5', 'Annually', '2025-01-01', 'active'],
            ]),
        })
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['created'], 1)
        self.assertEqual(response.data['updated'], 0)
        self.assertEqual(response.data['errors'], [])
        imported = AuditUniverse.objects.get(code=code)
        self.assertEqual(imported.name, 'Imported Process')
        self.assertEqual(imported.department, self.department)
        self.assertEqual(str(imported.risk_score), '3.50')

    def test_import_creates_from_xlsx(self):
        code = f'UNV-XLS-{next_seq()}'
        response = self.as_user(self.manager).post(self.IMPORT_URL, {
            'file': self.make_xlsx([
                self.CSV_HEADER.split(','),
                [code, 'Imported From Excel', 'system', '', '', '', 'Owner', 4, 'Quarterly', '2025-03-15', 'active'],
            ]),
        })
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['created'], 1)
        entry = AuditUniverse.objects.get(code=code)
        self.assertEqual(entry.category, 'system')
        self.assertEqual(entry.owner, 'Owner')

    def test_import_upserts_by_code_without_growing_the_table(self):
        # Columns: code, name, category, dept, dir, description, OWNER,
        # risk_score, frequency, last_audited, status.
        rows = [
            self.CSV_HEADER.split(','),
            [self.entry.code, self.entry.name, 'process', self.department.code, '',
             '', 'Renamed offline', '4.0', 'Annually', '', 'active'],
        ]
        before = AuditUniverse.objects.count()

        # A fresh SimpleUploadedFile per request — a reused one is already read
        # past its end and uploads as empty on the second call.
        first = self.as_user(self.manager).post(self.IMPORT_URL, {'file': self.make_csv(rows)})
        self.assertEqual(first.data['updated'], 1, first.data)
        self.entry.refresh_from_db()
        self.assertEqual(self.entry.owner, 'Renamed offline')
        self.assertEqual(str(self.entry.risk_score), '4.00')

        # Blank cells on the update leave the existing value alone (a lossless
        # export -> edit -> re-import cycle), and re-running is idempotent.
        second = self.as_user(self.manager).post(self.IMPORT_URL, {'file': self.make_csv(rows)})
        self.assertEqual(second.data['updated'], 0, second.data)

        self.assertEqual(AuditUniverse.objects.count(), before)
        self.entry.refresh_from_db()
        self.assertEqual(self.entry.owner, 'Renamed offline')

    def test_import_commits_valid_rows_and_reports_the_bad_ones(self):
        good_code = f'UNV-GOOD-{next_seq()}'
        response = self.as_user(self.manager).post(self.IMPORT_URL, {
            'file': self.make_csv([
                self.CSV_HEADER.split(','),
                [good_code, 'Good Row', 'process', self.department.code, '', '', '', '3', '', '', 'active'],
                [f'UNV-BAD-{next_seq()}', 'Bad Category', 'bogus', '', '', '', '', '3', '', '', 'active'],
                [f'UNV-BAD-{next_seq()}', 'Bad Department', 'process', 'NO-SUCH-DEPT', '', '', '', '3', '', '', 'active'],
            ]),
        })
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['created'], 1)
        self.assertTrue(AuditUniverse.objects.filter(code=good_code).exists())
        messages = {e['message'] for e in response.data['errors']}
        self.assertIn("invalid category 'bogus'", str(messages))
        self.assertTrue(any('unknown department_code' in m for m in messages))
        # Both bad rows must be reported with their spreadsheet row numbers
        # (header is row 1, so the bad rows sit at 3 and 4).
        self.assertEqual({e['row'] for e in response.data['errors']}, {3, 4})

    def test_import_rejects_missing_identity_columns(self):
        response = self.as_user(self.manager).post(self.IMPORT_URL, {
            'file': self.make_csv([['name', 'category'], ['Only A Name', 'process']]),
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn('code', response.data['detail'])

    def test_import_rejects_unhandled_file_type(self):
        upload = SimpleUploadedFile('universe.txt', b'code,name\n', content_type='text/plain')
        response = self.as_user(self.manager).post(self.IMPORT_URL, {'file': upload})
        self.assertEqual(response.status_code, 400)

    def test_import_requires_write_audit_export_stays_open(self):
        code = f'UNV-NOPE-{next_seq()}'
        response = self.as_user(self.auditee).post(self.IMPORT_URL, {
            'file': self.make_csv([
                self.CSV_HEADER.split(','),
                [code, 'Blocked', 'process', '', '', '', '', '3', '', '', 'active'],
            ]),
        })
        self.assertEqual(response.status_code, 403)
        self.assertFalse(AuditUniverse.objects.filter(code=code).exists())
        # Reads (export) are open to every authenticated role.
        self.assertEqual(self.as_user(self.auditee).get(self.EXPORT_URL).status_code, 200)

    def test_import_and_export_are_audit_logged(self):
        client = self.as_user(self.auditor)
        client.post(self.IMPORT_URL, {
            'file': self.make_csv([
                self.CSV_HEADER.split(','),
                [f'UNV-TRAIL-{next_seq()}', 'Trailed Import', 'process', self.department.code,
                 '', '', '', '', '', '', 'active'],
            ]),
        })
        client.get(self.EXPORT_URL)

        actions = list(
            AuditTrail.objects.filter(model_name='AuditUniverse').values_list('action', flat=True)
        )
        self.assertIn('IMPORT', actions)
        self.assertIn('EXPORT', actions)


class ProjectRegistryRoleAccessTest(RoleFixtureMixin, TestCase):
    """PPM project registry: WRITE_AUDIT gates writes; reads open to every role."""

    def test_every_role_can_list(self):
        Project.objects.create(code=f'PRJ-LIST-{next_seq()}', name='Listable Project')
        self.assert_status_by_role(
            {role: 200 for role in self.users},
            lambda client, role: client.get(PROJECTS_URL),
        )

    def test_only_write_audit_roles_can_create(self):
        def create(client, role):
            # Unique per role: five roles POSTing the same code would fail on the
            # unique constraint rather than on the permission.
            return client.post(PROJECTS_URL, {
                'code': f'PRJ-{role[:3].upper()}-{next_seq()}',
                'name': f'Project for {role}',
                'department': self.department.id,
            }, format='json')

        self.assert_status_by_role({
            Role.ADMIN: 201,
            Role.AUDIT_MANAGER: 201,
            Role.SUPERVISOR: 201,
            Role.AUDITOR: 201,
            Role.AUDITEE: 403,
        }, create)

    def test_auditee_post_does_not_persist(self):
        self.as_user(self.auditee).post(PROJECTS_URL, {
            'code': f'PRJ-AUD-{next_seq()}', 'name': 'No',
        }, format='json')
        self.assertFalse(Project.objects.filter(name='No').exists())

    def test_created_project_is_persisted_and_relistable(self):
        payload = {'code': f'PRJ-PERSIST-{next_seq()}', 'name': 'Persistent Project'}
        created = self.as_user(self.auditor).post(PROJECTS_URL, payload, format='json')
        self.assertEqual(created.status_code, 201)
        listed = self.as_user(self.auditor).get(PROJECTS_URL).data['results']
        codes = [p['code'] for p in listed]
        self.assertIn(payload['code'], codes)

    def test_duplicate_project_code_rejected(self):
        code = f'PRJ-DUP-{next_seq()}'
        Project.objects.create(code=code, name='First')
        response = self.as_user(self.auditor).post(PROJECTS_URL, {
            'code': code, 'name': 'Second',
        }, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Project.objects.filter(code=code).count(), 1)

    def test_project_create_is_audit_logged(self):
        self.as_user(self.auditor).post(PROJECTS_URL, {
            'code': f'PRJ-LOG-{next_seq()}', 'name': 'Logged',
        }, format='json')
        self.assertTrue(
            AuditTrail.objects.filter(model_name='Project', action='CREATE').exists()
        )


class DueForReAuditTest(RoleFixtureMixin, TestCase):
    """``due-for-re-audit`` — the Phase 3.3 re-audit window."""

    def setUp(self):
        super().setUp()
        today = timezone.now().date()
        self.never_audited = make_universe(
            department=self.department, audit_frequency='Annually', last_audited=None,
        )
        self.overdue = make_universe(
            department=self.department, audit_frequency='Annually',
            last_audited=today - datetime.timedelta(days=400),
        )
        self.recent = make_universe(
            department=self.department, audit_frequency='Annually',
            last_audited=today - datetime.timedelta(days=30),
        )
        self.system_overdue = make_universe(
            department=self.department, category='system', audit_frequency='Quarterly',
            last_audited=today - datetime.timedelta(days=200),
        )

    def codes(self, response):
        payload = response.data
        rows = payload['results'] if isinstance(payload, dict) else payload
        return {row['code'] for row in rows}

    def test_lists_never_audited_and_lapsed_entries_only(self):
        response = self.as_user(self.auditor).get(f'{UNIVERSE_URL}due-for-re-audit/')
        self.assertEqual(response.status_code, 200)
        codes = self.codes(response)
        self.assertIn(self.never_audited.code, codes)
        self.assertIn(self.overdue.code, codes)
        self.assertNotIn(self.recent.code, codes)

    def test_inactive_entries_are_excluded(self):
        retired = make_universe(
            department=self.department, audit_frequency='Annually',
            last_audited=None, status='inactive',
        )
        response = self.as_user(self.auditor).get(f'{UNIVERSE_URL}due-for-re-audit/')
        self.assertNotIn(retired.code, self.codes(response))

    def test_as_of_shifts_the_window(self):
        """A "what if" date in the past pulls the recently audited entry out of
        scope; the same entry is due again when asked about a date far ahead."""
        client = self.as_user(self.auditor)
        past = (self.recent.last_audited + datetime.timedelta(days=1)).isoformat()
        self.assertNotIn(
            self.recent.code,
            self.codes(client.get(f'{UNIVERSE_URL}due-for-re-audit/?as_of={past}')),
        )
        future = (self.recent.last_audited + datetime.timedelta(days=800)).isoformat()
        self.assertIn(
            self.recent.code,
            self.codes(client.get(f'{UNIVERSE_URL}due-for-re-audit/?as_of={future}')),
        )

    def test_category_narrows_the_list(self):
        response = self.as_user(self.auditor).get(
            f'{UNIVERSE_URL}due-for-re-audit/?category=system'
        )
        self.assertEqual(self.codes(response), {self.system_overdue.code})

    def test_bad_as_of_date_is_a_400_not_a_500(self):
        response = self.as_user(self.auditor).get(
            f'{UNIVERSE_URL}due-for-re-audit/?as_of=not-a-date'
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('YYYY-MM-DD', response.data['detail'])

    def test_response_is_paginated(self):
        response = self.as_user(self.auditor).get(f'{UNIVERSE_URL}due-for-re-audit/')
        self.assertIn('results', response.data)
        self.assertIn('count', response.data)


class AuditPlanWorkflowTest(RoleFixtureMixin, TestCase):
    """Submit is owner-or-approver; approve needs APPROVE_PLANS."""

    def setUp(self):
        super().setUp()
        self.plan = make_plan(created_by=self.auditor)

    def submit(self, user):
        return self.as_user(user).post(f'{PLANS_URL}{self.plan.id}/submit/')

    def test_author_can_submit_and_approvers_are_notified(self):
        response = self.submit(self.auditor)
        self.assertEqual(response.status_code, 200)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.status, 'submitted')

        for approver in (self.admin, self.manager):
            self.assertTrue(
                Notification.objects.filter(
                    user=approver, notification_type='approval_needed',
                ).exists(),
                f'{approver.role} was not asked to approve',
            )
        # The submitter is not told about their own submission, and a supervisor
        # is not on the plan-approval distribution list.
        self.assertEqual(notification_titles(self.auditor), [])
        self.assertEqual(notification_titles(self.supervisor), [])

    def test_an_uninvolved_auditor_cannot_submit_someone_elses_plan(self):
        """The gate this test exists for: at the class-level WRITE_AUDIT
        permission any auditor could push a colleague's draft to the approvers
        under their own name."""
        other_auditor = self.make_user(Role.AUDITOR, department=self.department)
        response = self.submit(other_auditor)
        self.assertEqual(response.status_code, 403)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.status, 'draft')

    def test_approve_plans_holders_can_submit_any_plan(self):
        self.assertEqual(self.submit(self.supervisor).status_code, 200)

    def test_auditee_cannot_submit(self):
        # 404 rather than 403 — see the note on the matching program test: the
        # plan is outside the auditee's scoped queryset, so it is refused before
        # the capability check rather than by it.
        self.assertEqual(self.submit(self.auditee).status_code, 404)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.status, 'draft')

    def test_submit_is_audit_logged_with_the_status_change(self):
        self.submit(self.auditor)
        entry = AuditTrail.objects.filter(
            model_name='AuditPlan', object_id=str(self.plan.id), action='UPDATE',
        ).first()
        self.assertIsNotNone(entry)
        self.assertEqual(entry.changes, {'status': ['draft', 'submitted']})

    def test_approve_stamps_the_approver_and_notifies_the_author(self):
        response = self.as_user(self.manager).post(f'{PLANS_URL}{self.plan.id}/approve/')
        self.assertEqual(response.status_code, 200)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.status, 'approved')
        self.assertEqual(self.plan.approved_by, self.manager)
        self.assertIsNotNone(self.plan.approved_at)
        self.assertTrue(
            Notification.objects.filter(
                user=self.auditor, notification_type='approved',
            ).exists()
        )

    def test_approve_is_gated_by_approve_plans(self):
        self.assert_status_by_role({
            Role.ADMIN: 200,
            Role.AUDIT_MANAGER: 200,
            Role.SUPERVISOR: 200,
            Role.AUDITOR: 403,
            Role.AUDITEE: 403,
        }, lambda client, role: client.post(f'{PLANS_URL}{self.plan.id}/approve/'))

    def test_approve_is_logged_as_an_approve_action(self):
        self.as_user(self.manager).post(f'{PLANS_URL}{self.plan.id}/approve/')
        self.assertTrue(
            AuditTrail.objects.filter(
                model_name='AuditPlan', object_id=str(self.plan.id), action='APPROVE',
            ).exists()
        )

    def test_create_records_the_author(self):
        response = self.as_user(self.auditor).post(PLANS_URL, {
            'title': 'FY Plan', 'year': timezone.now().year,
        }, format='json')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(
            AuditPlan.objects.get(pk=response.data['id']).created_by, self.auditor,
        )

    def test_auditee_cannot_create_a_plan(self):
        response = self.as_user(self.auditee).post(PLANS_URL, {
            'title': 'Blocked Plan', 'year': timezone.now().year,
        }, format='json')
        self.assertEqual(response.status_code, 403)


class AuditEngagementTest(RoleFixtureMixin, TestCase):
    """Engagement creation, team membership and status transitions."""

    def setUp(self):
        super().setUp()
        self.plan = make_plan(created_by=self.manager)

    def test_create_generates_a_number_and_notifies_the_team(self):
        response = self.as_user(self.manager).post(ENGAGEMENTS_URL, {
            'plan': self.plan.id,
            'title': 'Revenue Assurance Audit',
            'engagement_type': 'financial',
            'department': self.department.id,
            'lead_auditor': self.auditor.id,
            'supervisor': self.supervisor.id,
            **ENGAGEMENT_SCOPE,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)

        engagement = AuditEngagement.objects.get(pk=response.data['id'])
        self.assertRegex(engagement.engagement_number, r'^ENG-\d{4}-\d{4}$')
        for member in (self.auditor, self.supervisor):
            self.assertTrue(
                Notification.objects.filter(
                    user=member, notification_type='assigned',
                ).exists(),
                f'{member.role} was not told about the assignment',
            )
        # The creator is not notified of their own assignment.
        self.assertEqual(notification_titles(self.manager), [])

    def test_the_auditee_representative_round_trips(self):
        """The representative is what a finding inherits its auditee from, so it
        has to be writable here and readable back — a silent drop would look
        exactly like a manager who never filled the field in."""
        response = self.as_user(self.manager).post(ENGAGEMENTS_URL, {
            'plan': self.plan.id,
            'title': 'Engagement with a named representative',
            'engagement_type': 'financial',
            'department': self.department.id,
            'auditee': self.auditee.id,
            **ENGAGEMENT_SCOPE,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        engagement = AuditEngagement.objects.get(pk=response.data['id'])
        self.assertEqual(engagement.auditee, self.auditee)
        self.assertEqual(response.data['auditee'], self.auditee.id)
        self.assertEqual(response.data['auditee_name'], self.auditee.full_name)

    def test_an_engagement_without_a_representative_reports_none(self):
        """`None`, not "Unassigned" as the lead auditor and supervisor report.
        A blank here is the reason this engagement's findings will reach nobody,
        so the client has to be able to tell the difference and prompt."""
        response = self.as_user(self.manager).post(ENGAGEMENTS_URL, {
            'plan': self.plan.id, 'title': 'No representative named',
            **ENGAGEMENT_SCOPE,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertIsNone(response.data['auditee'])
        self.assertIsNone(response.data['auditee_name'])

    def test_client_supplied_engagement_number_is_ignored(self):
        response = self.as_user(self.manager).post(ENGAGEMENTS_URL, {
            'plan': self.plan.id, 'title': 'Spoofed Number',
            'engagement_number': 'ENG-HACKED',
            **ENGAGEMENT_SCOPE,
        }, format='json')
        self.assertEqual(response.status_code, 201)
        self.assertNotEqual(response.data['engagement_number'], 'ENG-HACKED')

    def test_auditee_cannot_create_an_engagement(self):
        response = self.as_user(self.auditee).post(ENGAGEMENTS_URL, {
            'plan': self.plan.id, 'title': 'Blocked Engagement',
        }, format='json')
        self.assertEqual(response.status_code, 403)

    def test_an_engagement_needs_both_an_objective_and_a_scope(self):
        """An engagement is where the audit is bounded — neither field is optional.

        Blank here does not mean "no constraints", it means an audit with no
        stated bounds, so the API refuses to create the row at all.
        """
        for omitted in ('objectives', 'scope'):
            with self.subTest(omitted=omitted):
                payload = {
                    'plan': self.plan.id,
                    'title': f'Unbounded Engagement ({omitted})',
                    **{k: v for k, v in ENGAGEMENT_SCOPE.items() if k != omitted},
                }
                response = self.as_user(self.manager).post(
                    ENGAGEMENTS_URL, payload, format='json',
                )
                self.assertEqual(response.status_code, 400, response.data)
                self.assertIn(omitted, response.data)

    def test_whitespace_only_scope_fields_are_refused(self):
        """A blank string is not a scope, so it must not slip past the check."""
        response = self.as_user(self.manager).post(ENGAGEMENTS_URL, {
            'plan': self.plan.id, 'title': 'Whitespace Engagement',
            'objectives': '   ', 'scope': '\n\t',
        }, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('objectives', response.data)
        self.assertIn('scope', response.data)

    def test_add_member_notifies_the_new_member(self):
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        second = self.make_user(Role.AUDITOR, department=self.department)
        response = self.as_user(self.supervisor).post(
            f'{ENGAGEMENTS_URL}{engagement.id}/add-member/',
            {'user': second.id, 'role': 'member', 'allocated_days': 5},
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(
            AuditTeamMember.objects.filter(engagement=engagement, user=second).exists()
        )
        self.assertTrue(
            Notification.objects.filter(user=second, notification_type='assigned').exists()
        )

    def test_update_status_stamps_the_actual_dates(self):
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        client = self.as_user(self.supervisor)

        client.post(f'{ENGAGEMENTS_URL}{engagement.id}/update-status/',
                    {'status': 'in_progress'}, format='json')
        engagement.refresh_from_db()
        self.assertEqual(engagement.status, 'in_progress')
        self.assertEqual(engagement.actual_start, timezone.now().date())

        client.post(f'{ENGAGEMENTS_URL}{engagement.id}/update-status/',
                    {'status': 'completed'}, format='json')
        engagement.refresh_from_db()
        self.assertEqual(engagement.actual_end, timezone.now().date())

    def test_completing_back_fills_the_linked_universe_entry(self):
        entry = make_universe(department=self.department, last_audited=None)
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, audit_universe=entry,
        )
        self.as_user(self.supervisor).post(
            f'{ENGAGEMENTS_URL}{engagement.id}/update-status/',
            {'status': 'completed'}, format='json',
        )
        entry.refresh_from_db()
        self.assertEqual(entry.last_audited, timezone.now().date())

    def test_completing_falls_back_to_the_department_universe_entry(self):
        """Engagements scheduled straight off a department, with no explicit
        universe link, must still close the re-audit loop — otherwise the
        entity looks permanently overdue."""
        low = make_universe(department=self.department, risk_score=1, last_audited=None)
        high = make_universe(department=self.department, risk_score=9, last_audited=None)
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor,
            audit_universe=None, department=self.department,
        )
        self.as_user(self.supervisor).post(
            f'{ENGAGEMENTS_URL}{engagement.id}/update-status/',
            {'status': 'completed'}, format='json',
        )
        low.refresh_from_db()
        high.refresh_from_db()
        # Highest risk score wins the fallback; the low-risk sibling is untouched.
        self.assertEqual(high.last_audited, timezone.now().date())
        self.assertIsNone(low.last_audited)

    def test_invalid_status_is_a_400(self):
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        response = self.as_user(self.supervisor).post(
            f'{ENGAGEMENTS_URL}{engagement.id}/update-status/',
            {'status': 'nonsense'}, format='json',
        )
        self.assertEqual(response.status_code, 400)
        engagement.refresh_from_db()
        self.assertEqual(engagement.status, 'planned')

    def test_reporting_notifies_supervisors_and_managers_but_not_the_actor(self):
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        self.as_user(self.supervisor).post(
            f'{ENGAGEMENTS_URL}{engagement.id}/update-status/',
            {'status': 'reporting'}, format='json',
        )
        self.assertTrue(notification_titles(self.manager))
        self.assertEqual(notification_titles(self.supervisor), [])

    # ── Completion guard ──────────────────────────────────────────────────
    # Completing an engagement stamps `actual_end` and back-fills
    # AuditUniverse.last_audited, which is what tells the re-audit due list the
    # entity has been covered. Before the guard, all of that could happen with
    # every finding still in draft. The four tests above use engagements with no
    # findings at all and must keep passing — they are the regression guard that
    # the new check only bites when there is something to bite on.

    def set_status(self, engagement, new_status):
        return self.as_user(self.supervisor).post(
            f'{ENGAGEMENTS_URL}{engagement.id}/update-status/',
            {'status': new_status}, format='json',
        )

    def test_an_open_finding_blocks_completion_and_writes_nothing(self):
        entry = make_universe(department=self.department, last_audited=None)
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, audit_universe=entry,
        )
        make_finding(engagement=engagement, identified_by=self.auditor, status='open')

        response = self.set_status(engagement, 'completed')

        self.assertEqual(response.status_code, 400, response.data)
        engagement.refresh_from_db()
        entry.refresh_from_db()
        # The refusal returns before the transaction opens, so none of the five
        # side effects of a completion happened.
        self.assertEqual(engagement.status, 'planned')
        self.assertIsNone(engagement.actual_end)
        self.assertIsNone(entry.last_audited)
        self.assertEqual(Notification.objects.count(), 0)
        self.assertFalse(
            AuditTrail.objects.filter(
                model_name='AuditEngagement', object_id=str(engagement.id),
            ).exists(),
            'a refused completion should not leave an audit-trail entry',
        )

    def test_the_refusal_names_the_blocking_findings(self):
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        blocker = make_finding(
            engagement=engagement, identified_by=self.auditor, status='open',
        )
        settled = make_finding(
            engagement=engagement, identified_by=self.auditor, status='closed',
        )

        response = self.set_status(engagement, 'completed')

        self.assertEqual(response.status_code, 400)
        detail = response.data['detail']
        self.assertIn(blocker.finding_number, detail)
        # Only the blockers are named — a settled finding is not the auditor's problem.
        self.assertNotIn(settled.finding_number, detail)
        self.assertIn('1 finding(s)', detail)

    def test_every_unsettled_status_blocks_completion(self):
        """`draft`, `open`, `in_progress` and `disputed` all block.

        Driven off BLOCKS_COMPLETION itself, so narrowing that constant
        narrows this test with it rather than leaving a stale assertion.
        """
        for finding_status in BLOCKS_COMPLETION:
            with self.subTest(finding_status=finding_status):
                engagement = make_engagement(
                    plan=self.plan, lead_auditor=self.auditor,
                )
                make_finding(
                    engagement=engagement, identified_by=self.auditor,
                    status=finding_status,
                )
                response = self.set_status(engagement, 'completed')
                self.assertEqual(response.status_code, 400, response.data)
                engagement.refresh_from_db()
                self.assertEqual(engagement.status, 'planned')

    def test_settled_findings_let_the_engagement_complete(self):
        entry = make_universe(department=self.department, last_audited=None)
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, audit_universe=entry,
        )
        make_finding(engagement=engagement, identified_by=self.auditor, status='resolved')
        make_finding(engagement=engagement, identified_by=self.auditor, status='closed')

        response = self.set_status(engagement, 'completed')

        self.assertEqual(response.status_code, 200, response.data)
        engagement.refresh_from_db()
        entry.refresh_from_db()
        self.assertEqual(engagement.status, 'completed')
        self.assertEqual(engagement.actual_end, timezone.now().date())
        self.assertEqual(entry.last_audited, timezone.now().date())

    def test_reporting_still_succeeds_with_an_open_finding(self):
        """The guard is completion-only — reporting is when findings are worked."""
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        make_finding(engagement=engagement, identified_by=self.auditor, status='open')

        response = self.set_status(engagement, 'reporting')

        self.assertEqual(response.status_code, 200, response.data)
        engagement.refresh_from_db()
        self.assertEqual(engagement.status, 'reporting')

    def test_cancelling_still_succeeds_with_an_open_finding(self):
        """Abandoning an engagement mid-flight is legitimate and must not
        require resolving the findings it already raised first."""
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        make_finding(engagement=engagement, identified_by=self.auditor, status='draft')

        response = self.set_status(engagement, 'cancelled')

        self.assertEqual(response.status_code, 200, response.data)
        engagement.refresh_from_db()
        self.assertEqual(engagement.status, 'cancelled')
        self.assertIsNone(engagement.actual_end)


class EngagementDirectorateInheritanceTest(RoleFixtureMixin, TestCase):
    """`directorate` is derived from the plan, never chosen.

    Every directorate-scoped number on the dashboard filters through
    ``AuditEngagement.directorate`` (``accounts.views.DashboardStatsView``), and
    nothing had ever written it: the engagement form has no such control and no
    signal filled it in, so ``?directorate=<id>`` matched nothing for all four
    directorates while the unscoped view looked healthy. The plan is now the
    single source of truth — these tests pin that, including the two ways the
    derivation could quietly go wrong again (a caller writing the field directly,
    and a plan edit leaving its engagements on the old directorate).
    """

    def setUp(self):
        super().setUp()
        self.fpa = Department.objects.create(
            name='Financial & Performance Audit', code=f'FPA-{next_seq()}',
            unit_type=Department.AUDIT, directorate_type='FPA',
        )
        self.ita = Department.objects.create(
            name='Information Technology Audit', code=f'ITA-{next_seq()}',
            unit_type=Department.AUDIT, directorate_type='ITA',
        )
        self.fpa_plan = make_plan(created_by=self.manager, directorate=self.fpa)
        self.ita_plan = make_plan(created_by=self.manager, directorate=self.ita)

    def test_a_new_engagement_inherits_its_plans_directorate(self):
        engagement = make_engagement(plan=self.fpa_plan, lead_auditor=self.auditor)
        self.assertEqual(engagement.directorate, self.fpa)

    def test_a_plan_without_a_directorate_leaves_the_engagement_without_one(self):
        """Strict inheritance. `make_plan` sets no directorate, and inventing one
        from the engagement's department would put a different answer in the
        column than the plan the badge claims to inherit from."""
        engagement = make_engagement(
            plan=make_plan(created_by=self.manager), lead_auditor=self.auditor,
        )
        self.assertIsNone(engagement.directorate)

    def test_moving_an_engagement_to_another_plan_re_derives(self):
        engagement = make_engagement(plan=self.fpa_plan, lead_auditor=self.auditor)
        engagement.plan = self.ita_plan
        engagement.save()
        engagement.refresh_from_db()
        self.assertEqual(engagement.directorate, self.ita)

    def test_a_partial_save_still_persists_the_derivation(self):
        """`update_fields` is how `update-status` saves. Without directorate in
        the field set the derivation would be computed and then dropped."""
        engagement = make_engagement(plan=self.fpa_plan, lead_auditor=self.auditor)
        engagement.plan = self.ita_plan
        engagement.save(update_fields=['plan'])
        engagement.refresh_from_db()
        self.assertEqual(engagement.directorate, self.ita)

    def test_the_api_ignores_a_client_supplied_directorate(self):
        """Read-only, so an engagement cannot be pinned to a directorate its plan
        does not belong to. Silently ignored rather than 400: the field is derived
        on every save, so there is nothing for the caller to correct."""
        response = self.as_user(self.manager).post(ENGAGEMENTS_URL, {
            'plan': self.fpa_plan.id,
            'title': 'Revenue Assurance Audit',
            'engagement_type': 'financial',
            'department': self.department.id,
            'lead_auditor': self.auditor.id,
            'directorate': self.ita.id,
            **ENGAGEMENT_SCOPE,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)

        engagement = AuditEngagement.objects.get(pk=response.data['id'])
        self.assertEqual(engagement.directorate, self.fpa)
        self.assertEqual(response.data['directorate'], self.fpa.id)

    def test_the_response_names_the_directorate_it_inherited(self):
        """The form's read-only badge renders these two fields, so they have to
        agree with the FK rather than being computed independently."""
        self.fpa.name_am = 'የፋይናንስ እና አፈጻጸም ኦዲት'
        self.fpa.save()
        response = self.as_user(self.manager).post(ENGAGEMENTS_URL, {
            'plan': self.fpa_plan.id,
            'title': 'Revenue Assurance Audit',
            'engagement_type': 'financial',
            'department': self.department.id,
            'lead_auditor': self.auditor.id,
            **ENGAGEMENT_SCOPE,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['directorate_name'], self.fpa.name)
        self.assertEqual(response.data['directorate_name_am'], self.fpa.name_am)

    def test_editing_a_plans_directorate_cascades_to_its_engagements(self):
        """The badge says "inherited from Plan". If the plan moves and its
        engagements do not, that sentence is a lie and the dashboard credits the
        directorate the plan no longer owns."""
        engagement = make_engagement(plan=self.fpa_plan, lead_auditor=self.auditor)
        self.fpa_plan.directorate = self.ita
        self.fpa_plan.save()

        engagement.refresh_from_db()
        self.assertEqual(engagement.directorate, self.ita)

    def test_clearing_a_plans_directorate_clears_its_engagements(self):
        engagement = make_engagement(plan=self.fpa_plan, lead_auditor=self.auditor)
        self.fpa_plan.directorate = None
        self.fpa_plan.save()

        engagement.refresh_from_db()
        self.assertIsNone(engagement.directorate)

    def test_a_plan_edit_that_leaves_the_directorate_alone_changes_nothing(self):
        """The cascade runs on every plan save. It must not disturb engagements
        when the directorate is untouched — or the approve/submit workflow, which
        saves the plan, would re-point its whole engagement list as a side effect."""
        engagement = make_engagement(plan=self.fpa_plan, lead_auditor=self.auditor)
        other = make_engagement(plan=self.ita_plan, lead_auditor=self.auditor)

        self.fpa_plan.title = 'Renamed Annual Plan'
        self.fpa_plan.save()

        engagement.refresh_from_db()
        other.refresh_from_db()
        self.assertEqual(engagement.directorate, self.fpa)
        self.assertEqual(other.directorate, self.ita)

    def test_the_cascade_reaches_engagements_it_never_saved_through_the_api(self):
        """Rows written straight to the ORM still follow the plan, because the
        rule lives on the model rather than in a serializer."""
        engagement = AuditEngagement.objects.create(
            plan=self.fpa_plan, title='Direct ORM write',
            engagement_number=f'ENG-DIRECT-{next_seq()}', objectives='o', scope='s',
        )
        self.assertEqual(engagement.directorate, self.fpa)

        self.fpa_plan.directorate = self.ita
        self.fpa_plan.save()
        engagement.refresh_from_db()
        self.assertEqual(engagement.directorate, self.ita)


class LifecycleStatusIsNotEditableTest(RoleFixtureMixin, TestCase):
    """`status` is the action routes' to move, not a field an edit may set.

    Every gate above lives in an `@action`: the engagement completion refusal, the
    plan's APPROVE_PLANS check, the stamps each approval writes. A PATCH that could
    write `status` reached past all of them, so the gates above were only ever
    enforced on the page that happened to use the action. These tests are the
    other half of those gates — they assert the side door is shut, which is the
    failure that leaves no trace (an engagement reads as closed, and nothing says
    it skipped the check).

    Read-only rather than refused: DRF drops the field and answers 200 with the
    record unchanged, so the assertion is on the stored row, not the status code.
    """

    def test_engagement_status_cannot_be_patched(self):
        """The one that matters: this skipped the open-findings completion guard.

        `update-status` refuses to complete an engagement while findings are
        unresolved. A PATCH could complete it anyway, and with it come the
        `actual_end` stamp and the universe's `last_audited` back-fill that tells
        the re-audit due list the entity has been covered.
        """
        engagement = make_engagement(lead_auditor=self.auditor)
        make_finding(engagement=engagement, identified_by=self.auditor, status='open')

        response = self.as_user(self.supervisor).patch(
            f'{ENGAGEMENTS_URL}{engagement.id}/',
            {'status': 'completed'}, format='json',
        )

        self.assertEqual(response.status_code, 200, response.data)
        engagement.refresh_from_db()
        self.assertEqual(engagement.status, 'planned')
        self.assertIsNone(engagement.actual_end)

    def test_engagement_actual_dates_cannot_be_backdated(self):
        """The stamps belong to `update-status` too, and are evidence of it.

        An auditor who writes `actual_end` by hand has asserted the engagement
        finished on a date the system never observed anything on.
        """
        engagement = make_engagement(lead_auditor=self.auditor)

        response = self.as_user(self.supervisor).patch(
            f'{ENGAGEMENTS_URL}{engagement.id}/',
            {
                'actual_start': (timezone.now().date() - datetime.timedelta(days=40)).isoformat(),
                'actual_end': (timezone.now().date() - datetime.timedelta(days=40)).isoformat(),
            },
            format='json',
        )

        self.assertEqual(response.status_code, 200, response.data)
        engagement.refresh_from_db()
        self.assertIsNone(engagement.actual_start)
        self.assertIsNone(engagement.actual_end)

    def test_plan_status_and_approver_cannot_be_patched(self):
        """A plan's approval is an APPROVE_PLANS sign-off that records who gave it."""
        plan = make_plan(created_by=self.auditor)

        response = self.as_user(self.auditor).patch(
            f'{PLANS_URL}{plan.id}/',
            {'status': 'approved', 'approved_by': self.manager.id}, format='json',
        )

        self.assertEqual(response.status_code, 200, response.data)
        plan.refresh_from_db()
        self.assertNotEqual(plan.status, 'approved')
        self.assertIsNone(plan.approved_by)
        self.assertIsNone(plan.approved_at)

    def test_plan_edits_still_work(self):
        """The positive control: only the lifecycle fields were closed off.

        Without this the tests above would also pass on a serializer that had
        stopped accepting any edits at all, which is not the change being made.
        """
        plan = make_plan(created_by=self.auditor)

        response = self.as_user(self.auditor).patch(
            f'{PLANS_URL}{plan.id}/',
            {'methodology': 'COSO 2013.'}, format='json',
        )

        self.assertEqual(response.status_code, 200, response.data)
        plan.refresh_from_db()
        self.assertEqual(plan.methodology, 'COSO 2013.')

    def test_engagement_edits_still_work(self):
        engagement = make_engagement(lead_auditor=self.auditor)

        response = self.as_user(self.auditor).patch(
            f'{ENGAGEMENTS_URL}{engagement.id}/',
            {'title': 'Renamed engagement'}, format='json',
        )

        self.assertEqual(response.status_code, 200, response.data)
        engagement.refresh_from_db()
        self.assertEqual(engagement.title, 'Renamed engagement')


class AuditEngagementScopingTest(RoleFixtureMixin, TestCase):
    """Auditees read only the engagements that concern their department."""

    def setUp(self):
        super().setUp()
        self.plan = make_plan(created_by=self.manager)
        self.own_dept = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, department=self.department,
        )
        self.other_dept = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, department=self.other_department,
        )

    def visible_ids(self, user):
        response = self.as_user(user).get(ENGAGEMENTS_URL)
        self.assertEqual(response.status_code, 200)
        return {row['id'] for row in response.data['results']}

    def test_auditee_sees_only_their_own_department(self):
        visible = self.visible_ids(self.auditee)
        self.assertIn(self.own_dept.id, visible)
        self.assertNotIn(self.other_dept.id, visible)

    def test_auditee_without_a_department_sees_only_named_engagements(self):
        """A missing department is read as "no department scope", not as
        "every department" — the safer reading of incomplete data."""
        floating = self.make_user(Role.AUDITEE, department=None)
        self.assertEqual(self.visible_ids(floating), set())

        AuditTeamMember.objects.create(
            engagement=self.other_dept, user=floating, role='member',
        )
        self.assertEqual(self.visible_ids(floating), {self.other_dept.id})

    def test_scoping_does_not_duplicate_rows(self):
        """An auditee who is both the department representative and a named team
        member matched two OR branches; without .distinct() the row came back
        twice and the count was wrong."""
        AuditTeamMember.objects.create(
            engagement=self.own_dept, user=self.auditee, role='member',
        )
        response = self.as_user(self.auditee).get(ENGAGEMENTS_URL)
        ids = [row['id'] for row in response.data['results']]
        self.assertEqual(ids, [self.own_dept.id])
        self.assertEqual(response.data['count'], 1)

    def test_other_roles_see_the_whole_calendar(self):
        for user in (self.admin, self.manager, self.supervisor, self.auditor):
            with self.subTest(role=user.role):
                visible = self.visible_ids(user)
                self.assertIn(self.own_dept.id, visible)
                self.assertIn(self.other_dept.id, visible)

    def test_auditee_cannot_retrieve_another_departments_engagement(self):
        response = self.as_user(self.auditee).get(
            f'{ENGAGEMENTS_URL}{self.other_dept.id}/'
        )
        self.assertEqual(response.status_code, 404)


class OrgScopeSplitTest(RoleFixtureMixin, TestCase):
    """Department, region and service center are three independent scopes.

    The picker used to store a single node of the corporate tree, so choosing a
    region or a service center *overwrote* the department — and because every
    region hangs off Region Coordination, picking Finance and then Adama Region
    left the department reading "Region Coordination". These pin the behaviour
    that replaced it: all three survive together.
    """

    def setUp(self):
        super().setUp()
        self.plan = make_plan(created_by=self.manager)
        self.coordination = Department.objects.create(
            name='Region Coordination', code='RGN Coordination',
            unit_type=Department.CORPORATE,
        )
        # Fixture department is "Finance <seq>" — the case from the report.
        self.region = Department.objects.create(
            name='Adama Region', code='RGN-BA',
            unit_type=Department.REGION, parent=self.coordination,
        )
        self.center = Department.objects.create(
            name='Adama CSC No.1', code='CSC-BA01',
            unit_type=Department.SERVICE_CENTER, parent=self.region,
        )

    def create_engagement(self, **overrides):
        payload = {
            'plan': self.plan.id,
            'title': 'Revenue Assurance Audit',
            'department': self.department.id,
            'region': self.region.id,
            'service_center': self.center.id,
            **ENGAGEMENT_SCOPE,
        }
        payload.update(overrides)
        return self.as_user(self.manager).post(ENGAGEMENTS_URL, payload, format='json')

    def test_a_region_does_not_replace_the_department(self):
        response = self.create_engagement()
        self.assertEqual(response.status_code, 201, response.data)
        # The point of the whole change: Finance is still the department.
        self.assertEqual(response.data['department'], self.department.id)
        self.assertEqual(response.data['region'], self.region.id)
        self.assertEqual(response.data['service_center'], self.center.id)

    def test_the_three_names_come_back_for_display(self):
        response = self.create_engagement()
        engagement = AuditEngagement.objects.get(pk=response.data['id'])
        self.assertEqual(engagement.department.name, self.department.name)
        self.assertEqual(response.data['department_name'], self.department.name)
        self.assertEqual(response.data['region_name'], 'Adama Region')
        self.assertEqual(response.data['service_center_name'], 'Adama CSC No.1')

    def test_department_alone_is_still_valid(self):
        """Region and service center are optional — most records set neither."""
        response = self.create_engagement(region=None, service_center=None)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertIsNone(response.data['region'])
        # A null relation must yield a null name, not a dropped key: the
        # tables read `*_name` directly.
        self.assertIsNone(response.data['region_name'])

    def test_clearing_the_region_leaves_the_department_alone(self):
        created = self.create_engagement()
        engagement_id = created.data['id']
        response = self.as_user(self.manager).patch(
            f'{ENGAGEMENTS_URL}{engagement_id}/',
            {'region': None, 'service_center': None}, format='json',
        )
        self.assertEqual(response.status_code, 200, response.data)
        engagement = AuditEngagement.objects.get(pk=engagement_id)
        self.assertIsNone(engagement.region)
        self.assertEqual(engagement.department_id, self.department.id)

    def test_the_scopes_are_independently_filterable(self):
        created = self.create_engagement()
        response = self.as_user(self.manager).get(
            ENGAGEMENTS_URL, {'region': self.region.id},
        )
        self.assertEqual(response.status_code, 200)
        ids = [row['id'] for row in response.data['results']]
        self.assertIn(created.data['id'], ids)


AUDITOR_PERFORMANCE_URL = f'{ENGAGEMENTS_URL}auditor-performance/'


class AuditorPerformanceTest(RoleFixtureMixin, TestCase):
    """Who carried what, and who finished the most.

    The endpoint unions three separate records — a person can be named on an
    engagement as `lead_auditor`, as `supervisor`, and as an `AuditTeamMember`
    — into one *set* of engagements per person. Most of these tests exist
    because adding those three counts together would look right on almost every
    row and be wrong on exactly the rows that matter, starting with the seeded
    directorate head who holds all three on every engagement.
    """

    def setUp(self):
        super().setUp()
        self.plan = make_plan(created_by=self.manager)

    def performance(self, user=None, **params):
        response = self.as_user(user or self.manager).get(
            AUDITOR_PERFORMANCE_URL, params,
        )
        self.assertEqual(response.status_code, 200, response.data)
        return response.data['auditors']

    @staticmethod
    def row_for(auditors, user):
        """The row for `user`, or None. Matched on id, so a shared display name
        cannot make one person's numbers stand in for another's."""
        return next((a for a in auditors if a['id'] == user.id), None)

    def add_member(self, engagement, user, role='member'):
        return AuditTeamMember.objects.create(
            engagement=engagement, user=user, role=role,
        )

    # ── The union ────────────────────────────────────────────────────────

    def test_the_whole_team_is_credited_not_only_the_lead(self):
        """A team member is credited for work the lead happens to be named on."""
        second = self.make_user(Role.AUDITOR, department=self.department)
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, status='completed',
        )
        self.add_member(engagement, second)

        auditors = self.performance()
        self.assertEqual(self.row_for(auditors, self.auditor)['completed'], 1)
        self.assertEqual(self.row_for(auditors, second)['completed'], 1)
        # ...but only the lead's `led` column moves, which is the whole reason
        # the two are reported separately.
        self.assertEqual(self.row_for(auditors, self.auditor)['led'], 1)
        self.assertEqual(self.row_for(auditors, second)['led'], 0)

    def test_one_person_named_three_times_is_credited_once(self):
        """The seeded shape: `lead_auditor` *and* `supervisor` *and* a team
        member — all the same person, on one engagement.

        Summing per-role counts gives 3 here. The set union gives 1, which is
        what the audit is worth to that person.
        """
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor,
            supervisor=self.auditor, status='completed',
        )
        self.add_member(engagement, self.auditor, role='lead')

        row = self.row_for(self.performance(), self.auditor)
        self.assertEqual(row['total'], 1)
        self.assertEqual(row['completed'], 1)
        self.assertEqual(row['led'], 1)

    def test_someone_with_no_engagements_is_not_ranked(self):
        make_engagement(plan=self.plan, lead_auditor=self.auditor)

        auditors = self.performance()
        self.assertIsNotNone(self.row_for(auditors, self.auditor))
        self.assertIsNone(self.row_for(auditors, self.auditee))

    def test_ranking_puts_the_most_completed_first(self):
        quiet = self.make_user(Role.AUDITOR, department=self.department)
        busy = self.make_user(Role.AUDITOR, department=self.department)
        make_engagement(plan=self.plan, lead_auditor=quiet, status='completed')
        for _ in range(3):
            make_engagement(plan=self.plan, lead_auditor=busy, status='completed')

        auditors = self.performance()
        order = [a['id'] for a in auditors]
        self.assertLess(order.index(busy.id), order.index(quiet.id))
        self.assertEqual(auditors[0]['completed'], 3)

    # ── Status buckets ───────────────────────────────────────────────────

    def test_cancelled_is_neither_completed_nor_in_flight(self):
        """A cancelled engagement is not output and it is not load. Counting it
        as either would misstate the person it was assigned to."""
        make_engagement(plan=self.plan, lead_auditor=self.auditor, status='cancelled')

        row = self.row_for(self.performance(), self.auditor)
        self.assertEqual(row['completed'], 0)
        self.assertEqual(row['in_flight'], 0)
        # It is still theirs, and still in the total — the row exists, it is
        # just not doing either job.
        self.assertEqual(row['total'], 1)
        self.assertEqual(row['by_status']['cancelled'], 1)

    def test_in_flight_covers_the_four_active_statuses(self):
        for engagement_status in ('planned', 'in_progress', 'fieldwork', 'reporting'):
            make_engagement(
                plan=self.plan, lead_auditor=self.auditor, status=engagement_status,
            )

        row = self.row_for(self.performance(), self.auditor)
        self.assertEqual(row['in_flight'], 4)
        self.assertEqual(row['completed'], 0)

    # ── On-time ──────────────────────────────────────────────────────────

    def test_on_time_ships_its_denominator(self):
        make_engagement(
            plan=self.plan, lead_auditor=self.auditor, status='completed',
            planned_end=datetime.date(2026, 6, 30),
            actual_end=datetime.date(2026, 6, 20),
        )
        make_engagement(
            plan=self.plan, lead_auditor=self.auditor, status='completed',
            planned_end=datetime.date(2026, 6, 30),
            actual_end=datetime.date(2026, 7, 15),
        )

        row = self.row_for(self.performance(), self.auditor)
        self.assertEqual(row['on_time'], {'met': 1, 'total': 2})
        self.assertEqual(row['on_time_rate'], 50)

    def test_an_engagement_with_no_planned_end_is_left_out_of_the_rate(self):
        """No deadline means nothing to have met. Counting it as on time would
        quietly inflate the score of whoever has the worst-planned work."""
        make_engagement(
            plan=self.plan, lead_auditor=self.auditor, status='completed',
            planned_end=None, actual_end=datetime.date(2026, 7, 15),
        )
        make_engagement(
            plan=self.plan, lead_auditor=self.auditor, status='completed',
            planned_end=datetime.date(2026, 6, 30),
            actual_end=datetime.date(2026, 7, 15),
        )

        row = self.row_for(self.performance(), self.auditor)
        self.assertEqual(row['on_time'], {'met': 0, 'total': 1})
        self.assertEqual(row['on_time_rate'], 0)

    def test_no_dated_completion_is_null_not_a_perfect_score(self):
        """The row must not read 100% for someone with nothing to measure.

        `completed` is 1 here and the rate is still None — the two columns are
        answering different questions.
        """
        make_engagement(
            plan=self.plan, lead_auditor=self.auditor, status='completed',
            planned_end=None, actual_end=None,
        )

        row = self.row_for(self.performance(), self.auditor)
        self.assertEqual(row['completed'], 1)
        self.assertEqual(row['on_time'], {'met': 0, 'total': 0})
        self.assertIsNone(row['on_time_rate'])

    # ── Findings and days ────────────────────────────────────────────────

    def test_findings_are_attributed_to_who_raised_them(self):
        """Two people on one engagement, each credited with their own findings.

        The supervisor has to be *named on the engagement* to have a row at all
        — that is the deliberate limit of this count. Findings are tallied
        inside a person's credited work, not across every finding they ever
        raised, so every column in a row describes the same set of audits. A
        finding raised on an engagement the person had no part in is not theirs
        to be ranked on.
        """
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, supervisor=self.supervisor,
        )
        make_finding(engagement=engagement, identified_by=self.auditor)
        make_finding(engagement=engagement, identified_by=self.auditor)
        make_finding(engagement=engagement, identified_by=self.supervisor)

        auditors = self.performance()
        self.assertEqual(self.row_for(auditors, self.auditor)['findings'], 2)
        self.assertEqual(self.row_for(auditors, self.supervisor)['findings'], 1)

    def test_a_finding_on_someone_elses_engagement_is_not_credited(self):
        """The other half of the rule above, asserted rather than implied: a
        person with no role on the engagement does not appear at all, however
        many findings they raised there."""
        engagement = make_engagement(plan=self.plan, lead_auditor=self.auditor)
        make_finding(engagement=engagement, identified_by=self.supervisor)

        self.assertIsNone(self.row_for(self.performance(), self.supervisor))

    def test_planned_days_sum_over_the_credited_engagements(self):
        make_engagement(plan=self.plan, lead_auditor=self.auditor, planned_days=12)
        make_engagement(plan=self.plan, lead_auditor=self.auditor, planned_days=30)

        row = self.row_for(self.performance(), self.auditor)
        self.assertEqual(row['planned_days'], 42)

    # ── The per-person engagement list ───────────────────────────────────

    def test_the_row_carries_its_own_engagements(self):
        """The expandable breakdown is served with the row, so opening one
        costs no extra request."""
        engagement = make_engagement(
            plan=self.plan, lead_auditor=self.auditor, status='completed',
        )

        row = self.row_for(self.performance(), self.auditor)
        listed = {e['id']: e for e in row['engagements']}
        self.assertIn(engagement.id, listed)
        self.assertEqual(listed[engagement.id]['number'], engagement.engagement_number)
        self.assertEqual(listed[engagement.id]['status'], 'completed')

    # ── Filters ──────────────────────────────────────────────────────────

    def test_year_narrows_to_that_plans_year(self):
        this_year = make_plan(created_by=self.manager, year=2026)
        last_year = make_plan(created_by=self.manager, year=2025)
        make_engagement(plan=this_year, lead_auditor=self.auditor, status='completed')
        make_engagement(plan=last_year, lead_auditor=self.auditor, status='completed')
        make_engagement(plan=last_year, lead_auditor=self.auditor, status='completed')

        self.assertEqual(self.row_for(self.performance(), self.auditor)['completed'], 3)
        narrowed = self.performance(year=2026)
        self.assertEqual(self.row_for(narrowed, self.auditor)['completed'], 1)

    def test_a_year_that_is_not_a_number_is_a_400(self):
        response = self.as_user(self.manager).get(
            AUDITOR_PERFORMANCE_URL, {'year': 'last year'},
        )
        self.assertEqual(response.status_code, 400)

    # ── Permission and scope ─────────────────────────────────────────────

    def test_the_auditee_is_refused(self):
        """The auditee holds no WRITE_AUDIT.

        Asserted directly rather than inferred from the capability matrix: this
        is a GET, and `HasCapability` lets *any* authenticated user through on a
        safe method — so a plain capability class would have leaked the ranking
        to the people being audited.
        """
        response = self.as_user(self.auditee).get(AUDITOR_PERFORMANCE_URL)
        self.assertEqual(response.status_code, 403)

    def test_every_audit_team_role_can_read_it(self):
        for user in (self.admin, self.manager, self.supervisor, self.auditor):
            with self.subTest(role=user.role):
                response = self.as_user(user).get(AUDITOR_PERFORMANCE_URL)
                self.assertEqual(response.status_code, 200)

    def test_the_ranking_honours_the_registers_region_scope(self):
        """It reads through `get_queryset()`, so a regional auditor's ranking is
        their region's work and not the whole organisation's.

        This is the test that fails if someone later swaps the call for a bare
        `AuditEngagement.objects` — a change that would look correct for every
        role except the one the scope exists for.
        """
        seq = next_seq()
        region_adama = Department.objects.create(
            name=f'Adama Region {seq}', code=f'RGN-ADAMA-{seq}',
            unit_type=Department.REGION,
        )
        region_hawassa = Department.objects.create(
            name=f'Hawassa Region {seq}', code=f'RGN-HAWASSA-{seq}',
            unit_type=Department.REGION,
        )
        adama_unit = Department.objects.create(
            name=f'Adama FPA Audit {seq}', code=f'FPA-RGN-ADAMA-{seq}',
            unit_type=Department.AUDIT, directorate_type='FPA', region=region_adama,
        )
        regional_auditor = self.make_user(Role.AUDITOR, department=adama_unit)

        make_engagement(
            plan=self.plan, lead_auditor=regional_auditor,
            region=region_adama, status='completed',
        )
        make_engagement(
            plan=self.plan, lead_auditor=regional_auditor,
            region=region_hawassa, status='completed',
        )

        # The manager carries no region on their department and sees both.
        self.assertEqual(
            self.row_for(self.performance(self.manager), regional_auditor)['completed'], 2,
        )
        # The regional auditor sees only their own region's engagement.
        row = self.row_for(self.performance(regional_auditor), regional_auditor)
        self.assertEqual(row['completed'], 1)
        self.assertEqual(row['total'], 1)
